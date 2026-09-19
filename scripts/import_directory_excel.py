"""Import SNHC Extension Directory Excel into seed JSON + Contact Center template."""

from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = Path(r"C:\Users\IT\Downloads\Directory Data (1).xlsx")
OUT_JSON = ROOT / "src" / "data" / "snhc-directory.json"
META_OUT = ROOT / "src" / "data" / "apps-script-directory-meta.json"
UPDATED_AT = "2026-09-15T00:00:00.000Z"
AS_ON = "15-Sep-2026"


def parse_staff_label(label: str) -> tuple[str, str, str]:
    """Return (person, designation, department)."""
    text = re.sub(r"\s+", " ", (label or "").strip())
    if not text:
        return "Unknown", "Staff", "—"

    if "," in text:
        left, right = text.rsplit(",", 1)
        person = left.strip(" -")
        dept = right.strip(" -")
        if person and dept and not re.fullmatch(r"\d+", dept):
            return person, "Staff", dept

    return text, "Staff", text


def cell_str(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def ext_str(value: object) -> str:
    s = cell_str(value)
    if not s:
        return ""
    # Excel may store extensions as numbers
    if re.fullmatch(r"\d+\.0", s):
        return s[:-2]
    return s


def pairs_from_row(row: tuple[object, ...], col_groups: list[tuple[int, int]]) -> list[tuple[str, str]]:
    out: list[tuple[str, str]] = []
    for name_i, ext_i in col_groups:
        name = cell_str(row[name_i] if name_i < len(row) else None)
        ext = ext_str(row[ext_i] if ext_i < len(row) else None)
        if name and ext:
            out.append((name, ext))
    return out


def main(xlsx: Path = DEFAULT_XLSX) -> None:
    if not xlsx.exists():
        raise SystemExit(f"Excel not found: {xlsx}")

    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=6, values_only=True))

    published = ""
    helpdesk_label = ""
    helpdesk_contact = ""
    section: str | None = None
    management: list[tuple[str, str]] = []
    staff: list[tuple[str, str]] = []

    for row in rows:
        a = cell_str(row[0] if row else None)
        e = cell_str(row[4] if len(row) > 4 else None)
        f = cell_str(row[5] if len(row) > 5 else None)

        if a == "Extension Directory" and e.lower().startswith("published"):
            published = f or e
            continue
        if a.startswith("IT Help Desk"):
            helpdesk_label = a
            helpdesk_contact = e or f
            continue
        if a == "Management Team":
            section = "mgmt"
            continue
        if a == "Departments & Staff":
            section = "staff"
            continue
        if a in {"Designation", "Name"} and "Extn" in cell_str(row[1] if len(row) > 1 else None):
            continue  # header row
        if not any(cell_str(c) for c in row):
            continue

        pairs = pairs_from_row(row, [(0, 1), (2, 3), (4, 5)])
        if not pairs:
            continue
        if section == "mgmt":
            management.extend(pairs)
        elif section == "staff":
            staff.extend(pairs)

    entries: list[dict] = []
    seen: set[str] = set()

    for i, (label, ext) in enumerate(management, start=1):
        key = f"mgmt|{ext}|{label.lower()}"
        if key in seen:
            continue
        seen.add(key)
        entries.append(
            {
                "id": f"as-mgmt-{i:03d}",
                "person": label,
                "designation": label,
                "department": "Management Team",
                "floor": "—",
                "zone": "—",
                "email": "",
                "mobile": "",
                "extension": ext,
                "status": "active",
                "updatedAt": UPDATED_AT,
            }
        )

    staff_i = 0
    for label, ext in staff:
        key = f"staff|{ext}|{label.lower()}"
        if key in seen:
            continue
        seen.add(key)
        staff_i += 1
        person, designation, department = parse_staff_label(label)
        entries.append(
            {
                "id": f"as-staff-{staff_i:03d}",
                "person": person,
                "designation": designation,
                "department": department,
                "floor": "—",
                "zone": "—",
                "email": "",
                "mobile": "",
                "extension": ext,
                "status": "active",
                "updatedAt": UPDATED_AT,
            }
        )

    # Parse helpdesk mobile / ext from "1559, 8800011659"
    hd_exts: list[str] = []
    hd_mobile = ""
    for part in re.split(r"[,/]", helpdesk_contact):
        part = part.strip()
        if not part:
            continue
        if len(part) >= 10:
            hd_mobile = part
        else:
            hd_exts.append(part)

    meta = {
        "source": "directory-excel",
        "file": xlsx.name,
        "publishedChip": f"Published {published}".strip(),
        "asOn": AS_ON,
        "mgmtCount": len(management),
        "staffCount": len(staff),
        "entryCount": len(entries),
        "helpdesk": {
            "label": helpdesk_label or "IT Help Desk (0930-1730) Week days",
            "extensions": hd_exts or ["1559"],
            "mobile": hd_mobile or "8800011659",
        },
    }

    OUT_JSON.write_text(json.dumps(entries, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    META_OUT.write_text(json.dumps(meta, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"Wrote {len(entries)} entries -> {OUT_JSON}")
    print(f"mgmt pairs {len(management)}, staff pairs {len(staff)}, unique {len(entries)}")
    print(f"asOn {AS_ON}, published {published}")
    print("sample", entries[0], entries[8] if len(entries) > 8 else None)


if __name__ == "__main__":
    import sys

    path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    main(path)
