"""Build SNHC Intercom Directory Excel for Contact Center team updates."""

from __future__ import annotations

import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src" / "data" / "snhc-directory.json"
OUT = ROOT / "public" / "templates" / "SNHC_Intercom_Directory_Template.xlsx"

HEADERS = [
    "Name",
    "Designation",
    "Department",
    "Section",
    "Floor",
    "Intercom",
    "Mobile",
    "Email",
    "Status",
]


def clean(value: object) -> str:
    s = "" if value is None else str(value).strip()
    if s in {"", "—", "–", "-", "N/A", "n/a", "Unknown"}:
        return ""
    return s


def main() -> None:
    rows = json.loads(SRC.read_text(encoding="utf-8"))
    sorted_rows = sorted(
        rows,
        key=lambda r: (clean(r.get("department")) or "ZZZ", clean(r.get("person")).lower()),
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Intercom Directory"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="00574E")
    thin = Border(
        left=Side(style="thin", color="D0D7D5"),
        right=Side(style="thin", color="D0D7D5"),
        top=Side(style="thin", color="D0D7D5"),
        bottom=Side(style="thin", color="D0D7D5"),
    )
    wrap = Alignment(vertical="center", wrap_text=True)

    for col, header in enumerate(HEADERS, 1):
        cell = ws.cell(1, col, header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", horizontal="left")
        cell.border = thin

    for i, row in enumerate(sorted_rows, 2):
        status = (
            "Inactive"
            if str(row.get("status") or "").lower() in {"inactive", "disabled"}
            else "Active"
        )
        values = [
            clean(row.get("person")),
            clean(row.get("designation")),
            clean(row.get("department")),
            clean(row.get("zone")),
            clean(row.get("floor")),
            clean(row.get("extension")),
            clean(row.get("mobile")),
            clean(row.get("email")),
            status,
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(i, col, val if val else None)
            cell.alignment = wrap
            cell.border = thin

    widths = [28, 28, 32, 16, 10, 12, 14, 36, 10]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.auto_filter.ref = f"A1:I{len(sorted_rows) + 1}"
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("Instructions", 0)
    ws2["A1"] = "SNHC Intercom Directory — Contact Center update sheet"
    ws2["A1"].font = Font(bold=True, size=14, color="00574E")
    ws2.merge_cells("A1:B1")

    notes = [
        "",
        "How to use",
        '1. Edit the "Intercom Directory" sheet — one row per contact.',
        "2. Fill blank Section, Floor, Mobile, and Email where known.",
        "3. Set Status to Active or Inactive (do not delete rows unless retiring a contact).",
        "4. Keep column headers exactly as-is for Admin import.",
        "5. Save as CSV (UTF-8) or Excel, then import via Contact Center Admin page.",
        "",
        "Columns",
        "Name — Person or desk name (required)",
        "Designation — Role / title",
        "Department — Department or location group",
        "Section — Zone / section within floor",
        "Floor — Floor label (G, 1, 2, …)",
        "Intercom — Extension number (required)",
        "Mobile — Direct mobile if available",
        "Email — Official email if available",
        "Status — Active or Inactive",
        "",
        f"Pre-filled rows: {len(sorted_rows)} contacts from SNHC Extension Directory (as on 15-Sep-2026).",
        "Blank Section / Floor / Mobile / Email — please fill location details and return to IT / Admin.",
    ]
    for i, line in enumerate(notes, 2):
        cell = ws2.cell(i, 1, line)
        if line in {"How to use", "Columns"}:
            cell.font = Font(bold=True, color="00574E")
    ws2.column_dimensions["A"].width = 92

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Wrote {OUT} with {len(sorted_rows)} contacts")


if __name__ == "__main__":
    main()
